import asyncio
import base64
import sys
import unittest
import io
from pathlib import Path
from urllib.parse import unquote
from unittest.mock import patch

AI_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(AI_ROOT))
sys.path.insert(0, str(AI_ROOT.parent / "scripts"))

from fastapi import HTTPException  # noqa: E402
from main import ProcessDocumentRequest, process_document  # noqa: E402
from docx import Document as WordDocument  # noqa: E402
from config import Settings  # noqa: E402
import smoke_document_lifecycle  # noqa: E402
from smoke_document_lifecycle import docx_bytes, pdf_bytes, validate_api_base  # noqa: E402

TEST_INGESTION_TOKEN = "test-only-document-ingestion-token-0123456789abcdef"


class DocumentIngestionTests(unittest.TestCase):
    def setUp(self):
        self.token_patch = patch("main.settings.DOCUMENT_INGESTION_TOKEN", TEST_INGESTION_TOKEN)
        self.token_patch.start()

    def tearDown(self):
        self.token_patch.stop()

    def test_smoke_never_sends_session_cookie_to_nonlocal_api(self):
        for remote in ("https://localhost:3001", "http://example.com", "http://localhost.evil.test", "http://user:pass@localhost:3001"):
            with self.subTest(api=remote), self.assertRaises(ValueError):
                validate_api_base(remote)
        self.assertEqual(validate_api_base("http://127.0.0.1:3001/"), "http://127.0.0.1:3001")

    def test_production_settings_reject_a_missing_secret(self):
        from pydantic import ValidationError
        with self.assertRaises(ValidationError):
            Settings(NODE_ENV="production", DOCUMENT_INGESTION_TOKEN="")

    def test_production_settings_accept_a_private_strong_token(self):
        settings = Settings(NODE_ENV="production", DOCUMENT_INGESTION_TOKEN=TEST_INGESTION_TOKEN)
        self.assertEqual(settings.DOCUMENT_INGESTION_TOKEN, TEST_INGESTION_TOKEN)

    def test_extracts_text_and_returns_project_scoped_vectors_without_persisting(self):
        request = ProcessDocumentRequest(
            document_id="60000000-0000-4000-8000-000000000001",
            project_id="60000000-0000-4000-8000-000000000002",
            filename="notes.txt",
            content_base64=base64.b64encode(b"Manual de acesso\n\nAutenticacao local").decode(),
        )
        async def embedding(text):
            self.assertTrue(text.strip())
            return [0.1] * 1024
        with patch("main.ollama_client.get_embedding", side_effect=embedding):
            result = asyncio.run(process_document(request, TEST_INGESTION_TOKEN))
        self.assertEqual(result["project_id"], request.project_id)
        self.assertEqual(len(result["chunks"]), 1)
        self.assertEqual(result["chunks"][0]["metadata"]["document_id"], request.document_id)
        self.assertEqual(len(result["chunks"][0]["embedding"]), 1024)

    def test_synthetic_smoke_fixtures_are_extractable_in_all_supported_formats(self):
        marker = "Quasar Nectario 7f6c-smoke-fixture"
        fixtures = (
            ("fixture.pdf", "application/pdf", pdf_bytes(marker)),
            ("fixture.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_bytes(marker)),
            ("fixture.md", "text/markdown", f"# Smoke\n\n{marker}\n".encode()),
            ("fixture.txt", "text/plain", f"Smoke documental.\n\n{marker}\n".encode()),
        )

        async def embedding(_text):
            return [0.1] * 1024

        for filename, _mime, content in fixtures:
            request = ProcessDocumentRequest(
                document_id="60000000-0000-4000-8000-000000000001",
                project_id="60000000-0000-4000-8000-000000000002",
                filename=filename,
                content_base64=base64.b64encode(content).decode(),
            )
            with self.subTest(filename=filename), patch("main.ollama_client.get_embedding", side_effect=embedding):
                result = asyncio.run(process_document(request, TEST_INGESTION_TOKEN))
                self.assertGreaterEqual(len(result["chunks"]), 1)
                self.assertIn(marker, "\n".join(chunk["text"] for chunk in result["chunks"]))
                self.assertTrue(all(len(chunk["embedding"]) == 1024 for chunk in result["chunks"]))

    def test_lifecycle_smoke_uploads_searches_and_removes_all_four_formats(self):
        project_id = "60000000-0000-4000-8000-000000000002"
        active = {}
        removed = set()

        def fake_request(url, _cookie, method="GET", body=None, headers=None):
            if method == "POST":
                document_id = f"60000000-0000-4000-8000-{len(active) + len(removed) + 1:012d}"
                filename = unquote(headers["X-File-Name"])
                active[document_id] = filename
                return 201, {"id": document_id, "nome": filename, "projeto_id": project_id}
            if method == "DELETE":
                document_id = url.rsplit("/", 1)[-1]
                active.pop(document_id, None)
                removed.add(document_id)
                return 204, None
            if "/documents?" in url:
                return 200, {"items": [
                    {"id": document_id, "projeto_id": project_id, "nome": filename, "status_processamento": "processado"}
                    for document_id, filename in active.items()
                ]}
            if "/search?" in url:
                return 200, {"items": [
                    {"entity_id": document_id, "project_id": project_id, "source_url": f"/projects/{project_id}/documents"}
                    for document_id in active if document_id not in removed
                ]}
            raise AssertionError(f"Rota inesperada no teste: {method} {url.split('?')[0]}")

        with patch.object(smoke_document_lifecycle, "request_json", side_effect=fake_request), patch("builtins.print"):
            smoke_document_lifecycle.run("http://127.0.0.1", project_id, "session-cookie-not-printed", 5)

        self.assertEqual(len(removed), 4)
        self.assertFalse(active)

    def test_docx_ingestion_extracts_paragraphs_and_table_cells_in_order(self):
        document = WordDocument()
        document.add_paragraph("Resumo do projeto")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Stack"
        table.cell(0, 1).text = "React"
        table.cell(1, 0).text = "Backend"
        table.cell(1, 1).text = "Node.js"
        document.add_paragraph("Decisão final")
        content = io.BytesIO()
        document.save(content)
        request = ProcessDocumentRequest(
            document_id="60000000-0000-4000-8000-000000000001",
            project_id="60000000-0000-4000-8000-000000000002",
            filename="architecture.docx",
            content_base64=base64.b64encode(content.getvalue()).decode(),
        )
        async def embedding(_text):
            return [0.1] * 1024
        with patch("main.ollama_client.get_embedding", side_effect=embedding):
            result = asyncio.run(process_document(request, TEST_INGESTION_TOKEN))
        extracted = result["chunks"][0]["text"]
        self.assertLess(extracted.index("Resumo do projeto"), extracted.index("Stack | React"))
        self.assertLess(extracted.index("Stack | React"), extracted.index("Backend | Node.js"))
        self.assertLess(extracted.index("Backend | Node.js"), extracted.index("Decisão final"))

    def test_rejects_empty_or_invalid_utf8_text(self):
        for content in (b"   ", b"\xff\xfe\xfa"):
            request = ProcessDocumentRequest(
                document_id="60000000-0000-4000-8000-000000000001",
                project_id="60000000-0000-4000-8000-000000000002",
                filename="notes.txt",
                content_base64=base64.b64encode(content).decode(),
            )
            with self.subTest(content=content), self.assertRaises(HTTPException) as raised:
                asyncio.run(process_document(request, TEST_INGESTION_TOKEN))
            self.assertEqual(raised.exception.status_code, 422)

    def test_rejects_more_than_500_chunks_before_calling_ollama(self):
        request = ProcessDocumentRequest(
            document_id="60000000-0000-4000-8000-000000000001",
            project_id="60000000-0000-4000-8000-000000000002",
            filename="large.txt",
            content_base64=base64.b64encode(b"x" * 501_000).decode(),
        )
        with patch("main.ollama_client.get_embedding") as embedding:
            with self.assertRaises(HTTPException) as raised:
                asyncio.run(process_document(request, TEST_INGESTION_TOKEN))
        self.assertEqual(raised.exception.status_code, 413)
        embedding.assert_not_called()

    def test_embedding_failure_does_not_return_partial_chunks(self):
        request = ProcessDocumentRequest(
            document_id="60000000-0000-4000-8000-000000000001",
            project_id="60000000-0000-4000-8000-000000000002",
            filename="notes.txt",
            content_base64=base64.b64encode(b"Manual de acesso").decode(),
        )
        async def unavailable(_text):
            raise RuntimeError("private ollama detail")
        with patch("main.ollama_client.get_embedding", side_effect=unavailable):
            with self.assertRaises(HTTPException) as raised:
                asyncio.run(process_document(request, TEST_INGESTION_TOKEN))
        self.assertEqual(raised.exception.status_code, 503)
        self.assertNotIn("private", raised.exception.detail)

    def test_rejects_request_without_valid_internal_token_before_ollama(self):
        request = ProcessDocumentRequest(
            document_id="60000000-0000-4000-8000-000000000001",
            project_id="60000000-0000-4000-8000-000000000002",
            filename="notes.txt",
            content_base64=base64.b64encode(b"content").decode(),
        )
        with patch("main.ollama_client.get_embedding") as embedding:
            with self.assertRaises(HTTPException) as raised:
                asyncio.run(process_document(request, "wrong-secret"))
        self.assertEqual(raised.exception.status_code, 401)
        embedding.assert_not_called()

    def test_rejects_ingestion_when_service_secret_is_not_configured(self):
        request = ProcessDocumentRequest(
            document_id="60000000-0000-4000-8000-000000000001",
            project_id="60000000-0000-4000-8000-000000000002",
            filename="notes.txt",
            content_base64=base64.b64encode(b"content").decode(),
        )
        with patch("main.settings.DOCUMENT_INGESTION_TOKEN", ""):
            with patch("main.ollama_client.get_embedding") as embedding:
                with self.assertRaises(HTTPException) as raised:
                    asyncio.run(process_document(request, TEST_INGESTION_TOKEN))
        self.assertEqual(raised.exception.status_code, 503)
        embedding.assert_not_called()


if __name__ == "__main__":
    unittest.main()
